import csv
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.models import ImportFileFormat, ImportLayout, ImportMappingInput

PREVIEW_ROW_LIMIT = 20
SUPPORTED_ENCODINGS = ("utf-8-sig", "gb18030")

FIELD_ALIASES = {
    "timestamp_column": (
        "timestamp",
        "source_timestamp",
        "datetime",
        "time",
        "时间",
        "采集时间",
        "数据时间",
    ),
    "tag_code_column": (
        "tag_code",
        "tag",
        "point_code",
        "point",
        "测点编码",
        "测点",
        "变量编码",
    ),
    "value_column": (
        "value",
        "numeric_value",
        "reading",
        "测量值",
        "数值",
        "值",
    ),
    "device_code_column": (
        "device_code",
        "device",
        "设备编码",
        "设备",
    ),
    "quality_column": (
        "quality",
        "status_code",
        "quality_code",
        "质量码",
        "状态码",
        "质量",
    ),
}


class TabularReadError(ValueError):
    pass


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _validate_headers(headers: list[str]) -> None:
    if not headers or all(not header for header in headers):
        raise TabularReadError("The file has no header row")
    if any(not header for header in headers):
        raise TabularReadError("Every source column must have a non-empty header")
    if len(set(headers)) != len(headers):
        raise TabularReadError("Source column headers must be unique")


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def detect_csv_encoding(path: Path) -> str:
    sample = path.read_bytes()[:65536]
    for encoding in SUPPORTED_ENCODINGS:
        try:
            sample.decode(encoding)
        except UnicodeDecodeError:
            continue
        return encoding
    raise TabularReadError("CSV must use UTF-8 or GB18030 encoding")


def _csv_dialect(path: Path, encoding: str) -> Any:
    with path.open("r", encoding=encoding, newline="") as source:
        sample = source.read(65536)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        return csv.excel


def iter_csv_rows(path: Path, encoding: str) -> Iterator[tuple[int, dict[str, Any]]]:
    dialect = _csv_dialect(path, encoding)
    with path.open("r", encoding=encoding, newline="") as source:
        reader = csv.DictReader(source, dialect=dialect)
        headers = [_normalize_header(value) for value in (reader.fieldnames or [])]
        _validate_headers(headers)
        reader.fieldnames = headers
        for row_number, row in enumerate(reader, start=2):
            yield (
                row_number,
                {key: value for key, value in row.items() if key is not None},
            )


def _xlsx_sheet(path: Path) -> tuple[Any, str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise TabularReadError("Unable to read the XLSX workbook") from exc
    worksheet = workbook.active
    if worksheet is None:
        workbook.close()
        raise TabularReadError("The XLSX workbook has no worksheet")
    return workbook, worksheet.title


def iter_xlsx_rows(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    workbook, sheet_name = _xlsx_sheet(path)
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        workbook.close()
        raise TabularReadError("The XLSX worksheet is empty") from exc
    headers = [_normalize_header(value) for value in header_row]
    _validate_headers(headers)
    try:
        for row_number, values in enumerate(rows, start=2):
            yield (
                row_number,
                {
                    header: _json_value(values[index] if index < len(values) else None)
                    for index, header in enumerate(headers)
                },
            )
    finally:
        workbook.close()


def iter_tabular_rows(
    path: Path,
    file_format: ImportFileFormat,
    encoding: str | None = None,
) -> Iterator[tuple[int, dict[str, Any]]]:
    if file_format == ImportFileFormat.CSV:
        yield from iter_csv_rows(path, encoding or detect_csv_encoding(path))
        return
    yield from iter_xlsx_rows(path)


def inspect_tabular_file(
    path: Path,
    file_format: ImportFileFormat,
) -> tuple[list[str], list[dict[str, Any]], str | None, str | None]:
    encoding = (
        detect_csv_encoding(path) if file_format == ImportFileFormat.CSV else None
    )
    rows = iter_tabular_rows(path, file_format, encoding)
    preview_rows: list[dict[str, Any]] = []
    source_columns: list[str] = []
    for _, row in rows:
        if not source_columns:
            source_columns = list(row)
        preview_rows.append(row)
        if len(preview_rows) >= PREVIEW_ROW_LIMIT:
            break
    if not source_columns:
        raise TabularReadError("The file contains a header but no data rows")
    sheet_name = None
    if file_format == ImportFileFormat.XLSX:
        workbook, sheet_name = _xlsx_sheet(path)
        workbook.close()
    return source_columns, preview_rows, encoding, sheet_name


def suggest_mapping(source_columns: list[str]) -> ImportMappingInput | None:
    normalized = {column.casefold(): column for column in source_columns}
    matches: dict[str, str | None] = {}
    for target, aliases in FIELD_ALIASES.items():
        matches[target] = next(
            (
                normalized[alias.casefold()]
                for alias in aliases
                if alias.casefold() in normalized
            ),
            None,
        )
    if not all(
        matches[field]
        for field in ("timestamp_column", "tag_code_column", "value_column")
    ):
        return None
    return ImportMappingInput(
        layout=ImportLayout.LONG,
        timestamp_column=str(matches["timestamp_column"]),
        tag_code_column=str(matches["tag_code_column"]),
        value_column=str(matches["value_column"]),
        device_code_column=matches["device_code_column"],
        quality_column=matches["quality_column"],
    )


def suggest_timestamp_column(source_columns: list[str]) -> str | None:
    normalized = {column.casefold(): column for column in source_columns}
    return next(
        (
            normalized[alias.casefold()]
            for alias in FIELD_ALIASES["timestamp_column"]
            if alias.casefold() in normalized
        ),
        None,
    )


def detect_layout(source_columns: list[str]) -> ImportLayout:
    return ImportLayout.LONG if suggest_mapping(source_columns) else ImportLayout.WIDE
